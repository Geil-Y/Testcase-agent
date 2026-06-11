"""Prompt Learning Requirement Excel parser."""

from __future__ import annotations

import io
from typing import Optional

from openpyxl import load_workbook

from .contracts import (
    RequirementColMap,
    ParsedRequirement,
    ParsedSummary,
    DataIssue,
)


def parse_requirements(
    file_bytes: bytes,
    sheet_name: str,
    mapping: RequirementColMap,
) -> tuple[list[ParsedRequirement], list[DataIssue]]:
    """Parse Requirement rows from the selected sheet using manual column mapping.

    Returns (parsed_requirements, data_issues).
    """
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet '{sheet_name}' not found in workbook")
    ws = wb[sheet_name]

    # Read header row
    headers: list[str] = []
    try:
        for cell in ws[1]:
            val = cell.value
            headers.append(str(val).strip() if val is not None else "")
    except IndexError:
        wb.close()
        return [], [DataIssue(
            file_type="requirement",
            sheet=sheet_name,
            row=None,
            issue_type="empty_sheet",
            message=f"Sheet '{sheet_name}' has no rows",
        )]

    # Resolve column indices
    key_col = _col_index(headers, mapping.requirement_key)
    desc_col = _col_index(headers, mapping.description)
    func_col = _col_index(headers, mapping.function_name) if mapping.function_name else -1
    type_col = _col_index(headers, mapping.requirement_type) if mapping.requirement_type else -1

    # Collect supplementary columns (all columns not otherwise mapped)
    mapped = {mapping.requirement_key, mapping.description}
    if mapping.function_name:
        mapped.add(mapping.function_name)
    if mapping.requirement_type:
        mapped.add(mapping.requirement_type)
    supp_cols = [(i, h) for i, h in enumerate(headers) if h and h not in mapped]

    requirements: list[ParsedRequirement] = []
    issues: list[DataIssue] = []
    seen_keys: set[str] = set()

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    for row_idx, row in enumerate(rows, start=2):
        key = _cell_str(row, key_col)
        desc = _cell_str(row, desc_col)

        # Missing requirement_key or description → exclude + data issue
        if not key or not desc:
            issues.append(DataIssue(
                file_type="requirement",
                sheet=sheet_name,
                row=row_idx,
                issue_type="missing_required_field",
                message=f"Row {row_idx}: {'missing requirement key' if not key else 'missing description'} — excluded",
            ))
            continue

        # Duplicate key → record issue, skip
        if key in seen_keys:
            issues.append(DataIssue(
                file_type="requirement",
                sheet=sheet_name,
                row=row_idx,
                issue_type="duplicate_key",
                message=f"Row {row_idx}: duplicate requirement key '{key}' — ignored",
            ))
            continue
        seen_keys.add(key)

        func = _cell_str(row, func_col)
        req_type = _cell_str(row, type_col) or "requirement"

        # Build supplementary_info dict from unmapped columns
        supplementary_info: dict[str, str] = {}
        for ci, ch in supp_cols:
            val = _cell_str(row, ci)
            if val:
                supplementary_info[ch] = val

        requirements.append(ParsedRequirement(
            requirement_key=key,
            description=desc,
            function_name=func,
            requirement_type=req_type,
            supplementary_info=supplementary_info,
            source_row=row_idx,
        ))

    return requirements, issues


def _col_index(headers: list[str], col_name: str) -> int:
    try:
        return headers.index(col_name)
    except ValueError:
        return -1


def _cell_str(row: tuple, index: int) -> str:
    if index < 0 or index >= len(row):
        return ""
    val = row[index]
    return str(val).strip() if val is not None else ""
